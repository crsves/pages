from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = Document()
def lvl(p, n):
    pPr = p._p.get_or_add_pPr(); numPr = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(n)); numPr.append(il)
    # reuse numId from style
    pPr.append(numPr)
p = d.add_paragraph("Emoji 🎉🎉 first then ")
r = p.add_run("BOLDWORD"); r.bold = True
p.add_run(" after.")
for t,n in [("lvl0 a",0),("lvl1 a",1),("lvl2 a",2),("lvl1 b",1),("lvl0 b",0)]:
    lvl(d.add_paragraph(t, style="List Number"), n)
p=d.add_paragraph("strike ")
r=p.add_run("gone"); r.font.strike=True
p.add_run(" H"); r=p.add_run("2"); r.font.subscript=True; p.add_run("O")
d.add_heading("Heading Three", 3)
d.save("fixture2.docx")
