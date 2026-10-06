from docx import Document
d = Document()
d.add_heading("The Fixture Essay", 0)
d.add_paragraph("A subtitle line", style="Subtitle")
d.add_heading("Introduction", 1)
p = d.add_paragraph("This paragraph has ")
p.add_run("bold").bold = True
p.add_run(", ")
p.add_run("italic").italic = True
p.add_run(", ")
r = p.add_run("underlined"); r.underline = True
p.add_run(", and ")
r = p.add_run("bold italic"); r.bold = True; r.italic = True
p.add_run(" text. " + "It is long enough to wrap across several lines in a narrow terminal window so that word wrapping logic gets exercised properly. " * 2)
d.add_heading("Lists", 2)
d.add_paragraph("First bullet", style="List Bullet")
d.add_paragraph("Second bullet", style="List Bullet")
d.add_paragraph("Nested bullet", style="List Bullet 2")
d.add_paragraph("Third bullet", style="List Bullet")
d.add_paragraph("Step one", style="List Number")
d.add_paragraph("Step two", style="List Number")
d.add_paragraph("Step three", style="List Number")
d.add_heading("A Table", 2)
t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
for i,row in enumerate([["Name","Role","Year"],["Ada","Analyst","1843"],["Grace","Admiral","1952"]]):
    for j,v in enumerate(row): t.cell(i,j).text = v
d.add_heading("Conclusion", 1)
d.add_paragraph("Final words with ünïcödé — and emoji 🎉 to test widths.")
q = d.add_paragraph("A block quote of some wisdom.", style="Quote")
d.save("fixture.docx")
