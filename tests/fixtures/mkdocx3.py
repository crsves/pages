from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx
d = Document()
p = d.add_paragraph("See ")
part = p.part
r_id = part.relate_to("https://example.com/docs", docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = "the docs"; r.append(t); h.append(r)
p._p.append(h)
p.add_run(" for more.")
t = d.add_table(rows=3, cols=2); t.style = "Table Grid"
trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), "true"); trPr.append(th)
for i,row in enumerate([["Item","Qty"],["Apples","12"],["Pears","3.5"]]):
    for j,v in enumerate(row): t.cell(i,j).text = v
d.add_page_break()
d.add_paragraph("After the page break.")
d.save("fixture3.docx")
