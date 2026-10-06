"""Build demo.docx; open it in Pages and save it as demo.pages.

Lists use explicit multi-level numbering definitions so Pages shows the
nesting (python-docx's built-in list styles import flat).
"""
import docx
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

d = Document()

BULLETS, NUMBERS = 91, 92


def define_lists():
    numbering = d.part.numbering_part.element

    def lvl(i, fmt, text):
        left = 360 * (i + 1)
        return (
            f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{left}" w:hanging="360"/></w:pPr></w:lvl>'
        )

    for num_id, levels in [
        (BULLETS, [lvl(0, "bullet", "•"), lvl(1, "bullet", "◦")]),
        (NUMBERS, [lvl(0, "decimal", "%1."), lvl(1, "lowerLetter", "%2.")]),
    ]:
        abstract = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{num_id}">'
            f'<w:multiLevelType w:val="hybridMultilevel"/>{"".join(levels)}</w:abstractNum>'
        )
        # abstractNum elements must come before any num element.
        first_num = numbering.find(qn("w:num"))
        first_num.addprevious(abstract) if first_num is not None else numbering.append(abstract)
        numbering.append(
            parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{num_id}"/></w:num>')
        )


def item(text, num_id, level):
    p = d.add_paragraph(text)
    num_pr = parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'
    )
    p._p.get_or_add_pPr().append(num_pr)


def link(p, text, url):
    r_id = p.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def runs(p, *parts):
    for part in parts:
        text, *fmt = part if isinstance(part, tuple) else (part,)
        r = p.add_run(text)
        r.bold = "b" in fmt
        r.italic = "i" in fmt
        r.underline = "u" in fmt
        r.font.strike = "s" in fmt
        r.font.superscript = "sup" in fmt
        r.font.subscript = "sub" in fmt
    return p


define_lists()

d.add_heading("The Case for Plain Text", 0)
d.add_paragraph("A short essay, written in Pages, read in a terminal", style="Subtitle")

d.add_heading("Why it matters", 1)
runs(
    d.add_paragraph(),
    "Documents outlive the apps that made them. A file you can ",
    ("read", "b"), ", ", ("search", "i"), ", and ", ("diff", "u"),
    " will still be useful in twenty years; one in a ",
    ("proprietary", "s"), " closed format might not be.",
)
p = runs(d.add_paragraph(), "That is why ", ("pages", "b"), " exists: ")
link(p, "github.com/crsves/pages", "https://github.com/crsves/pages")

d.add_heading("What a good format gives you", 2)
for text, level in [
    ("Longevity", 0),
    ("Readable without the original app", 1),
    ("No licence, account, or cloud required", 1),
    ("Tooling", 0),
    ("grep, diff, and version control just work", 1),
]:
    item(text, BULLETS, level)

d.add_heading("How to adopt it", 2)
for text, level in [
    ("Keep the canonical copy in a durable format", 0),
    ("Export from rich editors on every save", 0),
    ("Automate it with a hook", 1),
    ("Check the exports into git", 1),
    ("Review changes as text, not screenshots", 0),
]:
    item(text, NUMBERS, level)

d.add_heading("Formats over time", 2)
rows = [
    ("Format", "Introduced", "Opens today?"),
    ("ASCII text", "1963", "Everywhere"),
    ("WordStar", "1978", "With converters"),
    ("Markdown", "2004", "Everywhere"),
    ("Pages", "2005", "Mac, iPhone, and now your terminal"),
]
t = d.add_table(rows=len(rows), cols=3)
t.style = "Table Grid"
header = OxmlElement("w:tblHeader")
header.set(qn("w:val"), "true")
t.rows[0]._tr.get_or_add_trPr().append(header)
for i, row in enumerate(rows):
    for j, value in enumerate(row):
        t.cell(i, j).text = value

d.add_paragraph(
    "Write programs to handle text streams, because that is a universal interface. (Doug McIlroy)",
    style="Quote",
)
runs(d.add_paragraph(), "Plain text is the H", ("2", "sub"), "O of computing: you only notice it when it's gone.")

d.save("demo.docx")
